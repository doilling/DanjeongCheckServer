import io,json,sys,tempfile,unittest
from unittest.mock import patch
import time
from pathlib import Path
from openpyxl import Workbook,load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from server import create_app

class ServerTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.app=create_app(self.tmp.name);self.app.testing=True;self.admin=self.app.test_client();self.tokens={}
  code=(Path(self.tmp.name)/'setup-code.txt').read_text()
  r=self.call(self.admin,'setup',{'code':code,'id':'admin','name':'관리자','password':'adminpw'})
  self.assertEqual(r.status_code,201)
  self.login(self.admin,'admin','adminpw')
  for u in [dict(role='student',name='김하늘',g='1',c='1',n='1',home='담임'),dict(role='student',name='박서연',g='1',c='2',n='1',home='다른담임'),dict(id='t1',role='teacher',name='담임',g='1',c='1'),dict(id='t2',role='teacher',name='다른교사')]:
   u['password']='1234';self.assertEqual(self.call(self.admin,'users',u).status_code,201)
 def tearDown(self): self.tmp.cleanup()
 def call(self,c,path,data=None,method=None):
  headers={'Origin':'http://localhost','X-CSRF-Token':self.tokens.get(id(c),'')}
  return c.open('/danjeong/api/'+path,method=method or ('POST' if data is not None else 'GET'),json=data,headers=headers)
 def login(self,c,uid,pw='1234',remember=False):
  r=self.call(c,'login',{'id':uid,'password':pw,'remember':remember});self.assertEqual(r.status_code,200,r.json);self.tokens[id(c)]=r.json['csrf'];return r
 def client(self,uid,pw='1234'):
  c=self.app.test_client();self.login(c,uid,pw);return c
 def add_record(self,c=None,sid='1101',note='안내'):
  r=self.call(c or self.admin,'records',{'sid':sid,'item':'명찰 미착용','note':note});self.assertEqual(r.status_code,201,r.json)
  return self.call(c or self.admin,'state').json['records'][0]['id']
 def upload(self,path,raw,name):
  return self.admin.post('/danjeong/api/'+path,data={'file':(io.BytesIO(raw),name)},headers={'Origin':'http://localhost','X-CSRF-Token':self.tokens[id(self.admin)]})
 def workbook(self,rows):
  wb=Workbook();ws=wb.active
  for row in rows:ws.append(row)
  out=io.BytesIO();wb.save(out);return out.getvalue()
 def test_setup_is_protected_and_cannot_repeat(self):
  self.assertEqual(self.call(self.admin,'setup',{}).status_code,409)
  app=create_app(Path(self.tmp.name)/'empty');c=app.test_client();r=c.post('/danjeong/api/setup',json={'code':'wrong'},headers={'Origin':'http://localhost'});self.assertEqual(r.status_code,403)
 def test_auth_csrf_and_no_password_leak(self):
  anon=self.app.test_client();self.assertEqual(self.call(anon,'state').status_code,401)
  self.assertEqual(self.admin.post('/danjeong/api/users',json={},headers={'Origin':'http://evil'}).status_code,403)
  self.assertEqual(self.admin.post('/danjeong/api/users',json={},headers={'Origin':'http://localhost'}).status_code,403)
  data=self.call(self.admin,'state').json
  for row in data['users']:self.assertNotIn('password_hash',row);self.assertNotIn('pw',row)
 def test_student_isolation_and_completion(self):
  rid=self.add_record();other=self.add_record(sid='1201');s=self.client('1101')
  state=self.call(s,'state').json;self.assertEqual([u['id'] for u in state['users']],['1101']);self.assertEqual(len(state['records']),1)
  self.assertEqual(self.call(s,'users',{}).status_code,403)
  self.assertEqual(self.call(s,f'records/{other}/history').status_code,403)
  self.assertEqual(self.call(s,f'records/{rid}',{'note':'forged'},'PUT').status_code,403)
  self.assertEqual(self.call(s,f'records/{rid}',{'status':'done'},'PUT').status_code,200)
  history=self.call(s,f'records/{rid}/history').json['history'];self.assertEqual(history[-1]['event'],'student_complete')
 def test_teacher_edit_permission_and_audit(self):
  teacher=self.client('t1');other=self.client('t2');rid=self.add_record(teacher)
  self.assertEqual(self.call(other,f'records/{rid}',{'note':'other'},'PUT').status_code,403)
  self.assertEqual(self.call(teacher,f'records/{rid}',{'note':'수정','status':'review'},'PUT').status_code,200)
  history=self.call(teacher,f'records/{rid}/history').json['history'];self.assertEqual(len(history),2);self.assertIn('수정',history[-1]['after_json'])
  self.assertEqual(self.call(teacher,'report').status_code,403)
  self.assertEqual(self.call(other,f'records/{rid}',{},'DELETE').status_code,403)
  self.assertEqual(self.call(teacher,f'records/{rid}',{},'DELETE').status_code,200)
 def test_user_crud_preserves_password_and_last_admin(self):
  self.assertEqual(self.call(self.admin,'users/admin',{},'DELETE').status_code,400)
  self.assertEqual(self.call(self.admin,'users/admin',{'name':'관리자','role':'teacher'},'PUT').status_code,400)
  self.assertEqual(self.call(self.admin,'users/1101',{'name':'수정학생','role':'student','g':'1','c':'1','n':'1','home':'담임'},'PUT').status_code,200)
  s=self.client('1101');self.assertEqual(self.call(s,'session').json['user']['name'],'수정학생')
  self.add_record();self.assertEqual(self.call(self.admin,'users/1101',{},'DELETE').status_code,200)
  self.assertFalse(any(r['sid']=='1101' for r in self.call(self.admin,'state').json['records']))
 def test_reset_scoped_and_session_revoked(self):
  teacher=self.client('t1');s=self.client('1101');self.assertEqual(self.call(teacher,'users/1201/reset',{}).status_code,403)
  self.assertEqual(self.call(teacher,'users/1101/reset',{}).status_code,200)
  self.assertEqual(self.call(s,'state').status_code,401)
  self.assertEqual(len(self.call(self.admin,'state').json['resets']),1)
 def test_password_change_revokes_all_sessions(self):
  one=self.client('1101');two=self.client('1101');self.assertEqual(self.call(one,'password',{'old':'wrong','new':'newpass'}).status_code,403)
  self.assertEqual(self.call(one,'password',{'old':'1234','new':'newpass'}).status_code,200)
  self.assertEqual(self.call(two,'state').status_code,401);self.login(one,'1101','newpass')
 def test_admin_can_reset_teacher_but_teacher_cannot(self):
  teacher=self.client('t1');other=self.client('t2')
  self.assertEqual(self.call(teacher,'users/t2/reset',{}).status_code,403)
  self.assertEqual(self.call(self.admin,'password',{'old':'adminpw','new':'adminpw','defaultStudentPw':'schoolpw'}).status_code,200)
  self.login(self.admin,'admin','adminpw')
  self.assertEqual(self.call(self.admin,'users/t2/reset',{}).status_code,200)
  self.assertEqual(self.call(other,'state').status_code,401)
  self.login(other,'t2','schoolpw')
  self.assertEqual(self.call(self.admin,'state').json['resets'][0]['sid'],'t2')
 def test_remember_session_survives_server_restart(self):
  c=self.app.test_client();r=self.login(c,'1101',remember=True);self.assertIn('Max-Age=34560000',r.headers['Set-Cookie']);self.assertIn('HttpOnly',r.headers['Set-Cookie'])
  token=c.get_cookie('danjeong_session',path='/danjeong/').value
  app=create_app(self.tmp.name);new=app.test_client();new.set_cookie('danjeong_session',token,path='/danjeong/')
  self.assertEqual(new.get('/danjeong/api/session').json['user']['id'],'1101')
 def test_login_limit_and_https_cookie(self):
  c=self.app.test_client()
  for _ in range(8): self.assertEqual(self.call(c,'login',{'id':'1101','password':'wrong'}).status_code,401)
  self.assertEqual(self.call(c,'login',{'id':'1101','password':'1234'}).status_code,429)
  with self.app.db() as db: db.execute('DELETE FROM login_attempts')
  r=c.post('/danjeong/api/login',json={'id':'1101','password':'1234'},headers={'Origin':'https://localhost','X-Forwarded-Proto':'https'})
  self.assertEqual(r.status_code,200,r.json);self.assertIn('Secure',r.headers['Set-Cookie'])
 def test_persistent_session_has_no_server_deadline_and_logout_revokes(self):
  c=self.app.test_client();self.login(c,'1101',remember=True)
  with self.app.db() as db:
   self.assertEqual(db.execute("SELECT expires FROM sessions WHERE uid='1101'").fetchone()[0],0)
  with patch('server.time.time',return_value=time.time()+500*86400):
   result=self.call(c,'state');self.assertEqual(result.status_code,200)
   self.assertIn('Max-Age=34560000',result.headers['Set-Cookie'])
   another=self.app.test_client();self.login(another,'t1')
   self.assertEqual(self.call(c,'state').status_code,200)
   self.assertEqual(self.call(c,'logout',{}).status_code,200)
   self.assertEqual(self.call(c,'state').status_code,401)
 def test_xlsx_atomic_import_and_real_export(self):
  raw=self.workbook([['학년','반','번호','이름','초기비밀번호'],[2,3,1,'새학생','abcd'],[2,3,2,'','abcd']])
  r=self.upload('import/students',raw,'students.xlsx');self.assertEqual(r.status_code,400);self.assertFalse(any(u['id']=='2301' for u in self.call(self.admin,'state').json['users']))
  raw=self.workbook([['학년','반','번호','이름','초기비밀번호'],[2,3,1,'새학생','abcd'],[1,1,1,'김하늘수정','']])
  r=self.upload('import/students',raw,'students.xlsx');self.assertEqual(r.status_code,200,r.json);self.assertEqual(r.json['created'],1);self.assertEqual(r.json['updated'],1)
  self.client('2301','abcd');self.client('1101')
  rid=self.add_record(note='=HYPERLINK("evil")')
  out=self.call(self.admin,'report.xlsx');self.assertEqual(out.status_code,200);wb=load_workbook(io.BytesIO(out.data));self.assertEqual(wb.active['H2'].data_type,'s');self.assertIn('HYPERLINK',wb.active['H2'].value)
  self.assertEqual(self.call(self.admin,'template/teachers.xlsx').status_code,200)
 def test_teacher_excel_and_dynamic_classes(self):
  raw=self.workbook([['교사아이디','이름','역할','담당학년','담당반','초기비밀번호'],['t3','새교사','교사',2,10,'abcd']])
  r=self.upload('import/teachers',raw,'teachers.xlsx');self.assertEqual(r.status_code,200,r.json);self.client('t3','abcd')
  self.assertEqual(self.call(self.admin,'users',dict(role='student',name='10반',g='2',c='10',n='1')).status_code,201)
  self.assertTrue(any(u['id']=='21001' for u in self.call(self.admin,'state').json['users']))
 def test_report_filters_repeat_and_xlsx(self):
  for _ in range(3):self.add_record()
  self.add_record(sid='1201');r=self.call(self.admin,'report?g=1&c=1&repeat=1&threshold=3').json['rows'];self.assertEqual(len(r),1);self.assertEqual(r[0]['count'],3)
  self.assertEqual(len(self.call(self.admin,'report?term=박').json['rows']),1)
  self.assertEqual(self.call(self.admin,'report?threshold=bad&repeat=1').status_code,400)
 def test_item_rules_crud(self):
  for table in ('items','rules'):
   self.assertEqual(self.call(self.admin,table,{'value':'새 안내'}).status_code,201)
   rid=self.call(self.admin,'state').json[table][-1]['id'];self.assertEqual(self.call(self.admin,f'{table}/{rid}',{'value':'수정 안내'},'PUT').status_code,200)
   self.assertEqual(self.call(self.admin,f'{table}/{rid}',{},'DELETE').status_code,200)
 def test_stale_record_edit_is_rejected(self):
  rid=self.add_record()
  self.assertEqual(self.call(self.admin,f'records/{rid}',{'note':'첫 수정','expectedVersion':1},'PUT').status_code,200)
  self.assertEqual(self.call(self.admin,f'records/{rid}',{'note':'오래된 화면','expectedVersion':1},'PUT').status_code,409)
  self.assertEqual(self.call(self.admin,'state').json['records'][0]['note'],'첫 수정')
 def test_deleted_items_and_rules_stay_deleted_after_restart(self):
  state=self.call(self.admin,'state').json
  for table in ('items','rules'):
   for row in state[table]: self.assertEqual(self.call(self.admin,f'{table}/{row["id"]}',{},'DELETE').status_code,200)
  restarted=create_app(self.tmp.name)
  with restarted.db() as db:
   self.assertEqual(db.execute('SELECT COUNT(*) FROM items').fetchone()[0],0)
   self.assertEqual(db.execute('SELECT COUNT(*) FROM rules').fetchone()[0],0)
 def test_backup_restore_is_transactional(self):
  self.add_record();backup=self.call(self.admin,'backup').data
  self.assertEqual(self.upload('restore',b'{}','bad.json').status_code,400)
  self.assertEqual(len(self.call(self.admin,'state').json['records']),1)
  corrupt=json.loads(backup);corrupt['records'].append(corrupt['records'][0])
  self.assertEqual(self.upload('restore',json.dumps(corrupt).encode(),'bad.json').status_code,400)
  self.assertEqual(len(self.call(self.admin,'state').json['records']),1)
  self.assertEqual(self.call(self.admin,'users',dict(role='student',name='추가',g='3',c='1',n='1')).status_code,201)
  self.assertEqual(self.upload('restore',backup,'backup.json').status_code,200)
  self.assertEqual(self.call(self.admin,'state').status_code,401);self.login(self.admin,'admin','adminpw')
  self.assertFalse(any(u['id']=='3101' for u in self.call(self.admin,'state').json['users']))
  self.assertTrue(list((Path(self.tmp.name)/'backups').glob('*.json')))
 def test_legacy_import_and_year_reset(self):
  legacy={'users':[{'id':'boss','name':'관리자','role':'admin','pw':'abcd'},{'id':'1101','name':'학생','role':'student','pw':'1234','g':'1','c':'1','n':'1'}],'items':['명찰'],'records':[{'id':10,'sid':'1101','item':'명찰','note':'안내','status':'pending','teacher':'선생님','date':'2026-10-01'}],'defaultStudentPw':'1234'}
  self.assertEqual(self.upload('restore',json.dumps(legacy).encode(),'legacy.json').status_code,200);self.login(self.admin,'boss','abcd')
  self.assertEqual(len(self.call(self.admin,'state').json['records']),1)
  self.assertEqual(self.call(self.admin,'year-reset',{'password':'wrong'}).status_code,403)
  self.assertEqual(self.call(self.admin,'year-reset',{'password':'abcd'}).status_code,200)
  data=self.call(self.admin,'state').json;self.assertEqual([u['id'] for u in data['users']],['boss']);self.assertEqual(data['records'],[])

if __name__=='__main__':unittest.main()
