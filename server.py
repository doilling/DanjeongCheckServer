"""Danjeong Check: central SQLite server, Windows/Waitress/PyInstaller."""
from __future__ import annotations
import argparse, csv, hashlib, hmac, io, json, logging, os, re, secrets, sqlite3, sys, time, zipfile
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from functools import wraps
from pathlib import Path
from flask import Flask, request, jsonify, g, send_file, send_from_directory
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent
BASE = Path(sys.executable).resolve().parent if getattr(sys, 'frozen', False) else ROOT
KST = timezone(timedelta(hours=9))
PREFIX = '/danjeong'
DEFAULT_ITEMS = ['명찰 미착용', '실내화 미착용', '교복 상의 규정 미준수']
DEFAULT_RULES = ['학교 지정 교복과 명찰을 단정하게 착용합니다.', '교내에서는 지정 실내화를 착용합니다.']
SCHEMA = '''
CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY,name TEXT NOT NULL,role TEXT NOT NULL CHECK(role IN ('student','teacher','admin')),password_hash TEXT NOT NULL,g TEXT NOT NULL DEFAULT '',c TEXT NOT NULL DEFAULT '',n TEXT NOT NULL DEFAULT '',home TEXT NOT NULL DEFAULT '');
CREATE TABLE IF NOT EXISTS items(id INTEGER PRIMARY KEY AUTOINCREMENT,value TEXT NOT NULL UNIQUE);
CREATE TABLE IF NOT EXISTS rules(id INTEGER PRIMARY KEY AUTOINCREMENT,value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS records(id INTEGER PRIMARY KEY AUTOINCREMENT,sid TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,item TEXT NOT NULL,note TEXT NOT NULL,status TEXT NOT NULL CHECK(status IN ('pending','review','done')),teacher_id TEXT NOT NULL,teacher TEXT NOT NULL,date TEXT NOT NULL,updated_at TEXT NOT NULL,version INTEGER NOT NULL DEFAULT 1);
CREATE INDEX IF NOT EXISTS records_student_date ON records(sid,date);
CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY AUTOINCREMENT,record_id INTEGER,event TEXT NOT NULL,actor TEXT NOT NULL,at TEXT NOT NULL,before_json TEXT,after_json TEXT);
CREATE TABLE IF NOT EXISTS resets(id INTEGER PRIMARY KEY AUTOINCREMENT,sid TEXT NOT NULL,name TEXT NOT NULL,processor TEXT NOT NULL,date TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY,uid TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,csrf TEXT NOT NULL,expires REAL NOT NULL);
CREATE TABLE IF NOT EXISTS login_attempts(identity TEXT PRIMARY KEY,count INTEGER NOT NULL,expires REAL NOT NULL);
'''

def stamp(): return datetime.now(KST).isoformat(timespec='seconds')
def today(): return datetime.now(KST).date().isoformat()
def token_hash(token): return hashlib.sha256(token.encode()).hexdigest()

class Problem(Exception):
    def __init__(self, message, status=400): self.message,self.status=message,status

def text(value, maxlen=1000):
    if not isinstance(value, (str,int)): raise Problem('입력 형식을 확인하세요.')
    value=str(value).strip()
    if len(value)>maxlen: raise Problem('입력 내용이 너무 깁니다.')
    return value

def password(value):
    if not isinstance(value,str) or not 4<=len(value)<=128: raise Problem('비밀번호는 4~128자로 입력하세요.')
    return value

def valid_date(value):
    try:
        if datetime.strptime(value,'%Y-%m-%d').strftime('%Y-%m-%d') != value: raise ValueError()
    except (ValueError,TypeError): raise Problem('날짜 형식을 확인하세요.')
    return value

def student_id(g,c,n):
    try:
        a,b,d=int(g),int(c),int(n)
        if not (1<=a<=9 and 1<=b<=99 and 1<=d<=99): raise ValueError()
    except (ValueError,TypeError): raise Problem('학년 1~9, 반·번호 1~99를 입력하세요.')
    return str(a)+str(b)+str(d).zfill(2)

def public_user(row): return {k:row[k] for k in ('id','name','role','g','c','n','home')}


def create_app(data_dir=None):
    app=Flask(__name__,static_folder=None)
    app.config.update(DATA=Path(data_dir or os.environ.get('DANJEONG_DATA',BASE/'data')),MAX_CONTENT_LENGTH=12*1024*1024)
    app.config['DATA'].mkdir(parents=True,exist_ok=True)
    app.config['DB']=app.config['DATA']/'danjeong.sqlite3'
    first_database=not app.config['DB'].exists()
    app.wsgi_app=ProxyFix(app.wsgi_app,x_for=1,x_proto=1,x_host=1)

    @contextmanager
    def db():
        conn=sqlite3.connect(app.config['DB'],timeout=20)
        conn.row_factory=sqlite3.Row
        conn.execute('PRAGMA foreign_keys=ON')
        try:
            with conn: yield conn
        finally: conn.close()
    app.db=db
    with db() as conn:
        conn.execute('PRAGMA journal_mode=WAL');conn.executescript(SCHEMA)
        conn.execute("INSERT OR IGNORE INTO settings VALUES('defaultStudentPw','1234')")
        if first_database:
            conn.executemany('INSERT INTO items(value) VALUES(?)',[(x,) for x in DEFAULT_ITEMS])
            conn.executemany('INSERT INTO rules(value) VALUES(?)',[(x,) for x in DEFAULT_RULES])
        initialized=bool(conn.execute('SELECT 1 FROM users').fetchone())
    setup_file=app.config['DATA']/'setup-code.txt'
    if not initialized and not setup_file.exists(): setup_file.write_text(secrets.token_urlsafe(24),encoding='utf-8')

    def body():
        data=request.get_json(silent=True)
        if not isinstance(data,dict): raise Problem('입력 형식을 확인하세요.')
        return data
    def settings(conn): return dict(conn.execute('SELECT key,value FROM settings'))
    def audit(conn,event,record_id=None,before=None,after=None):
        conn.execute('INSERT INTO audit(record_id,event,actor,at,before_json,after_json) VALUES(?,?,?,?,?,?)',
            (record_id,event,g.user['id'] if getattr(g,'user',None) else 'setup',stamp(),json.dumps(before,ensure_ascii=False) if before else None,json.dumps(after,ensure_ascii=False) if after else None))
    def auth(*roles):
        def decorate(fn):
            @wraps(fn)
            def wrapped(*args,**kwargs):
                if not getattr(g,'user',None): raise Problem('로그인이 필요합니다.',401)
                if roles and g.user['role'] not in roles: raise Problem('접근 권한이 없습니다.',403)
                if request.method not in ('GET','HEAD') and not hmac.compare_digest(request.headers.get('X-CSRF-Token',''),g.csrf): raise Problem('요청을 확인할 수 없습니다. 새로고침 후 다시 시도하세요.',403)
                return fn(*args,**kwargs)
            return wrapped
        return decorate
    @app.before_request
    def protect():
        g.user=None;g.csrf=None
        if request.method not in ('GET','HEAD','OPTIONS'):
            origin=request.headers.get('Origin')
            if origin != request.host_url.rstrip('/'): raise Problem('요청 출처가 올바르지 않습니다.',403)
        token=request.cookies.get('danjeong_session','')
        if token:
            with db() as conn:
                row=conn.execute('SELECT s.csrf,u.* FROM sessions s JOIN users u ON u.id=s.uid WHERE s.token_hash=? AND s.expires>?',(token_hash(token),time.time())).fetchone()
            if row: g.user=public_user(row);g.csrf=row['csrf']
    @app.after_request
    def headers(response):
        response.headers.update({'Cache-Control':'no-store','X-Content-Type-Options':'nosniff','X-Frame-Options':'DENY','Referrer-Policy':'same-origin','Content-Security-Policy':"default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; base-uri 'none'; frame-ancestors 'none'; form-action 'self'"})
        return response
    @app.errorhandler(Problem)
    def problem(ex): return jsonify(error=ex.message),ex.status
    @app.errorhandler(sqlite3.IntegrityError)
    def conflict(ex): return jsonify(error='중복된 아이디·항목 또는 연결된 자료를 확인하세요.'),409
    @app.errorhandler(413)
    def large(ex): return jsonify(error='업로드 파일은 12MB 이하로 선택하세요.'),413
    @app.errorhandler(404)
    def missing(ex): return jsonify(error='페이지를 찾을 수 없습니다.'),404
    @app.errorhandler(Exception)
    def unexpected(ex):
        app.logger.exception('Request failed')
        return jsonify(error='처리 중 오류가 발생했습니다. 서버 로그를 확인하세요.'),500

    @app.get(PREFIX)
    def redirect_root():
        from flask import redirect
        return redirect(PREFIX+'/')
    @app.get(PREFIX+'/')
    def index(): return send_from_directory(ROOT/'static','index.html')
    @app.get(PREFIX+'/<path:path>')
    def static(path):
        if path.startswith('api/'): raise Problem('페이지를 찾을 수 없습니다.',404)
        return send_from_directory(ROOT/'static',path)
    @app.get(PREFIX+'/api/health')
    def health(): return jsonify(ok=True,version='5.0.0')
    @app.get(PREFIX+'/api/session')
    def session():
        with db() as conn: initialized=bool(conn.execute('SELECT 1 FROM users').fetchone())
        return jsonify(user=g.user,csrf=g.csrf,initialized=initialized)
    @app.post(PREFIX+'/api/setup')
    def setup():
        d=body()
        with db() as conn:
            if conn.execute('SELECT 1 FROM users').fetchone(): raise Problem('초기 설정이 이미 완료되었습니다.',409)
            code=d.get('code','')
            if not isinstance(code,str) or not setup_file.exists() or not hmac.compare_digest(code,setup_file.read_text().strip()): raise Problem('서버의 setup-code.txt에 있는 설정 코드를 확인하세요.',403)
            uid=text(d.get('id','admin'),50);name=text(d.get('name','관리자'),100)
            if not re.fullmatch(r'[A-Za-z0-9_-]{1,50}',uid) or not name: raise Problem('관리자 아이디와 이름을 확인하세요.')
            conn.execute('INSERT INTO users(id,name,role,password_hash) VALUES(?,?,?,?)',(uid,name,'admin',generate_password_hash(password(d.get('password')))))
            audit(conn,'initial_setup')
        setup_file.unlink(missing_ok=True)
        return jsonify(message='초기 관리자 설정 완료'),201
    @app.post(PREFIX+'/api/login')
    def login():
        d=body();uid=text(d.get('id',''),50);pw=d.get('password','');identity=uid+'|'+request.remote_addr
        with db() as conn:
            attempt=conn.execute('SELECT * FROM login_attempts WHERE identity=?',(identity,)).fetchone()
            if attempt and attempt['expires']>time.time() and attempt['count']>=8: raise Problem('로그인 시도가 많습니다. 15분 후 다시 시도하세요.',429)
            row=conn.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone()
            ok=row and isinstance(pw,str) and len(pw)<=128 and check_password_hash(row['password_hash'],pw)
        if not ok:
            with db() as conn:
                current=time.time()
                conn.execute('INSERT INTO login_attempts VALUES(?,1,?) ON CONFLICT(identity) DO UPDATE SET count=CASE WHEN login_attempts.expires>? THEN login_attempts.count+1 ELSE 1 END,expires=excluded.expires',(identity,current+900,current))
            raise Problem('아이디 또는 비밀번호를 확인하세요.',401)
        token=secrets.token_urlsafe(32);csrf=secrets.token_urlsafe(24);remember=d.get('remember') is True;duration=30*86400 if remember else 8*3600
        with db() as conn:
            conn.execute('DELETE FROM login_attempts WHERE identity=?',(identity,));conn.execute('DELETE FROM sessions WHERE expires<?',(time.time(),))
            conn.execute('INSERT INTO sessions VALUES(?,?,?,?)',(token_hash(token),uid,csrf,time.time()+duration))
        response=jsonify(user=public_user(row),csrf=csrf)
        response.set_cookie('danjeong_session',token,httponly=True,secure=request.is_secure,samesite='Strict',path=PREFIX+'/',max_age=duration if remember else None)
        return response
    @app.post(PREFIX+'/api/logout')
    @auth()
    def logout():
        with db() as conn: conn.execute('DELETE FROM sessions WHERE token_hash=?',(token_hash(request.cookies.get('danjeong_session','')),))
        response=jsonify(message='로그아웃 완료');response.delete_cookie('danjeong_session',path=PREFIX+'/');return response
    @app.get(PREFIX+'/api/state')
    @auth()
    def state():
        with db() as conn:
            student=g.user['role']=='student'
            users=[public_user(r) for r in conn.execute('SELECT * FROM users '+('WHERE id=?' if student else '')+' ORDER BY role,id',(g.user['id'],) if student else ())]
            records=[dict(r) for r in conn.execute('SELECT * FROM records '+('WHERE sid=?' if student else '')+' ORDER BY date DESC,id DESC',(g.user['id'],) if student else ())]
            values=settings(conn)
            resets=[dict(r) for r in conn.execute('SELECT * FROM resets ORDER BY id DESC')] if g.user['role']=='admin' else []
            return jsonify(users=users,records=records,items=[dict(r) for r in conn.execute('SELECT * FROM items ORDER BY id')],rules=[dict(r) for r in conn.execute('SELECT * FROM rules ORDER BY id')],defaultStudentPw=values['defaultStudentPw'] if g.user['role']=='admin' else None,resets=resets,today=today())

    def make_user(d,old=None):
        role=d.get('role',old['role'] if old else 'student')
        if role not in ('student','teacher','admin'): raise Problem('역할을 확인하세요.')
        name=text(d.get('name',''),100)
        if not name: raise Problem('이름을 입력하세요.')
        grade,cls,number=[text(d.get(k,''),10) for k in ('g','c','n')]
        if role=='student': generated=student_id(grade,cls,number);grade,cls,number=map(lambda x:str(int(x)),(grade,cls,number))
        else: generated=''
        uid=old['id'] if old else (generated if role=='student' else text(d.get('id',''),50))
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,50}',uid): raise Problem('아이디는 영문·숫자·밑줄·하이픈으로 입력하세요.')
        if old and old['role']!=role and ('student' in (old['role'],role)): raise Problem('학생과 교직원 계정 유형은 변경할 수 없습니다.')
        if old and old['role']=='admin' and role!='admin':
            with db() as conn:
                if old['id']==g.user['id'] or conn.execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0]<=1: raise Problem('현재 관리자 또는 마지막 관리자의 권한은 변경할 수 없습니다.')
        pw=d.get('password')
        hashed=generate_password_hash(password(pw)) if pw else (old['password_hash'] if old else None)
        if hashed is None:
            with db() as conn: hashed=generate_password_hash(password(settings(conn)['defaultStudentPw'] if role=='student' else '1234'))
        return dict(id=uid,name=name,role=role,password_hash=hashed,g=grade,c=cls,n=number,home=text(d.get('home',''),100))
    def write_user(conn,u):
        keys=('id','name','role','password_hash','g','c','n','home')
        conn.execute('INSERT INTO users VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,role=excluded.role,password_hash=excluded.password_hash,g=excluded.g,c=excluded.c,n=excluded.n,home=excluded.home',[u[k] for k in keys])
    @app.post(PREFIX+'/api/users')
    @auth('admin')
    def add_user():
        u=make_user(body())
        with db() as conn:
            if conn.execute('SELECT 1 FROM users WHERE id=?',(u['id'],)).fetchone(): raise Problem('이미 등록된 아이디입니다.',409)
            write_user(conn,u);audit(conn,'user_create',after=public_user(u))
        return jsonify(message='계정을 등록했습니다.',user=public_user(u)),201
    @app.put(PREFIX+'/api/users/<uid>')
    @auth('admin')
    def update_user(uid):
        with db() as conn: old=conn.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone()
        if not old: raise Problem('계정을 찾을 수 없습니다.',404)
        u=make_user(body(),old)
        with db() as conn:
            write_user(conn,u);audit(conn,'user_update',before=public_user(old),after=public_user(u))
            if u['password_hash']!=old['password_hash'] or u['role']!=old['role']: conn.execute('DELETE FROM sessions WHERE uid=?',(uid,))
        return jsonify(message='계정을 수정했습니다.')
    @app.delete(PREFIX+'/api/users/<uid>')
    @auth('admin')
    def delete_user(uid):
        with db() as conn:
            row=conn.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone()
            if not row: raise Problem('계정을 찾을 수 없습니다.',404)
            if uid==g.user['id'] or (row['role']=='admin' and conn.execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0]<=1): raise Problem('현재 관리자 또는 마지막 관리자는 삭제할 수 없습니다.')
            audit(conn,'user_delete',before=public_user(row));conn.execute('DELETE FROM users WHERE id=?',(uid,))
        return jsonify(message='계정을 삭제했습니다.')
    @app.post(PREFIX+'/api/password')
    @auth()
    def change_password():
        d=body();new=password(d.get('new'))
        with db() as conn:
            row=conn.execute('SELECT * FROM users WHERE id=?',(g.user['id'],)).fetchone()
            if not isinstance(d.get('old'),str) or not check_password_hash(row['password_hash'],d['old']): raise Problem('현재 비밀번호가 일치하지 않습니다.',403)
            conn.execute('UPDATE users SET password_hash=? WHERE id=?',(generate_password_hash(new),g.user['id']))
            if g.user['role']=='admin' and 'defaultStudentPw' in d: conn.execute("UPDATE settings SET value=? WHERE key='defaultStudentPw'",(password(d['defaultStudentPw']),))
            conn.execute('DELETE FROM sessions WHERE uid=?',(g.user['id'],));audit(conn,'password_change')
        return jsonify(message='비밀번호가 변경되었습니다. 다시 로그인하세요.')
    @app.post(PREFIX+'/api/users/<uid>/reset')
    @auth('admin','teacher')
    def reset_password(uid):
        with db() as conn:
            row=conn.execute("SELECT * FROM users WHERE id=? AND role='student'",(uid,)).fetchone()
            if not row: raise Problem('학생을 찾을 수 없습니다.',404)
            if g.user['role']=='teacher' and not ((g.user['g'] and g.user['c'] and row['g']==g.user['g'] and row['c']==g.user['c']) or row['home']==g.user['name']): raise Problem('담당 학반 또는 담임 학생만 초기화할 수 있습니다.',403)
            conn.execute('UPDATE users SET password_hash=? WHERE id=?',(generate_password_hash(settings(conn)['defaultStudentPw']),uid));conn.execute('DELETE FROM sessions WHERE uid=?',(uid,))
            conn.execute('INSERT INTO resets(sid,name,processor,date) VALUES(?,?,?,?)',(uid,row['name'],g.user['name'],stamp()));audit(conn,'password_reset',after={'sid':uid})
        return jsonify(message='학교 기본 비밀번호로 초기화했습니다.')
    @app.route(PREFIX+'/api/<collection>',methods=['POST'])
    @auth('admin')
    def add_value(collection):
        if collection not in ('items','rules'): raise Problem('경로를 확인하세요.',404)
        value=text(body().get('value',''),500)
        if not value: raise Problem('내용을 입력하세요.')
        with db() as conn: conn.execute(f'INSERT INTO {collection}(value) VALUES(?)',(value,));audit(conn,collection+'_create',after={'value':value})
        return jsonify(message='추가했습니다.'),201
    @app.route(PREFIX+'/api/<collection>/<int:rid>',methods=['PUT','DELETE'])
    @auth('admin')
    def edit_value(collection,rid):
        if collection not in ('items','rules'): raise Problem('경로를 확인하세요.',404)
        with db() as conn:
            if not conn.execute(f'SELECT 1 FROM {collection} WHERE id=?',(rid,)).fetchone(): raise Problem('항목을 찾을 수 없습니다.',404)
            if request.method=='DELETE': conn.execute(f'DELETE FROM {collection} WHERE id=?',(rid,))
            else:
                value=text(body().get('value',''),500)
                if not value: raise Problem('내용을 입력하세요.')
                conn.execute(f'UPDATE {collection} SET value=? WHERE id=?',(value,rid))
            audit(conn,collection+'_update',after={'id':rid})
        return jsonify(message='저장했습니다.')

    @app.post(PREFIX+'/api/records')
    @auth('admin','teacher')
    def add_record():
        d=body();sid=text(d.get('sid',''),50);item=text(d.get('item',''),500);note=text(d.get('note',''),4000)
        with db() as conn:
            if not conn.execute("SELECT 1 FROM users WHERE id=? AND role='student'",(sid,)).fetchone(): raise Problem('등록된 학생을 선택하세요.')
            if not conn.execute('SELECT 1 FROM items WHERE value=?',(item,)).fetchone(): raise Problem('등록된 점검 항목을 선택하세요.')
            rid=conn.execute('INSERT INTO records(sid,item,note,status,teacher_id,teacher,date,updated_at) VALUES(?,?,?,?,?,?,?,?)',(sid,item,note,'pending',g.user['id'],g.user['name'],today(),stamp())).lastrowid
            audit(conn,'record_create',rid,after=dict(conn.execute('SELECT * FROM records WHERE id=?',(rid,)).fetchone()))
        return jsonify(message='지도 내역을 등록했습니다.'),201
    def get_record(conn,rid):
        r=conn.execute('SELECT * FROM records WHERE id=?',(rid,)).fetchone()
        if not r: raise Problem('기록을 찾을 수 없습니다.',404)
        if g.user['role']=='student' and r['sid']!=g.user['id']: raise Problem('접근 권한이 없습니다.',403)
        return r
    @app.get(PREFIX+'/api/records/<int:rid>/history')
    @auth()
    def history(rid):
        with db() as conn:
            get_record(conn,rid)
            return jsonify(history=[dict(r) for r in conn.execute('SELECT * FROM audit WHERE record_id=? ORDER BY id',(rid,))])
    @app.put(PREFIX+'/api/records/<int:rid>')
    @auth()
    def edit_record(rid):
        d=body()
        with db() as conn:
            r=get_record(conn,rid)
            if 'expectedVersion' in d and d['expectedVersion']!=r['version']: raise Problem('다른 사용자가 이 기록을 변경했습니다. 화면을 새로 조회한 뒤 수정하세요.',409)
            if g.user['role']=='student':
                if not set(d)<= {'status','expectedVersion'} or d.get('status')!='done': raise Problem('학생은 본인의 교정 완료만 처리할 수 있습니다.',403)
                item,note,date,status=r['item'],r['note'],r['date'],'done'
            else:
                if g.user['role']=='teacher' and r['teacher_id']!=g.user['id']: raise Problem('본인이 등록한 기록만 수정할 수 있습니다.',403)
                item=text(d.get('item',r['item']),500);note=text(d.get('note',r['note']),4000);date=valid_date(d.get('date',r['date']));status=d.get('status',r['status'])
                if not item or status not in ('pending','review','done'): raise Problem('점검 항목·처리 상태를 확인하세요.')
            conn.execute('UPDATE records SET item=?,note=?,date=?,status=?,updated_at=?,version=version+1 WHERE id=?',(item,note,date,status,stamp(),rid))
            audit(conn,'student_complete' if g.user['role']=='student' else 'record_update',rid,dict(r),dict(conn.execute('SELECT * FROM records WHERE id=?',(rid,)).fetchone()))
        return jsonify(message='지도 내역을 저장했습니다.')
    @app.delete(PREFIX+'/api/records/<int:rid>')
    @auth('admin','teacher')
    def delete_record(rid):
        with db() as conn:
            r=get_record(conn,rid)
            if g.user['role']=='teacher' and r['teacher_id']!=g.user['id']: raise Problem('본인이 등록한 기록만 삭제할 수 있습니다.',403)
            audit(conn,'record_delete',rid,before=dict(r));conn.execute('DELETE FROM records WHERE id=?',(rid,))
        return jsonify(message='지도 내역을 삭제했습니다.')

    def filtered(conn):
        clauses=[];args=[]
        for key,column in [('from','r.date'),('to','r.date')]:
            val=request.args.get(key)
            if val: clauses.append(column+('>=?' if key=='from' else '<=?'));args.append(valid_date(val))
        for key in ('g','c'):
            val=request.args.get(key)
            if val: clauses.append('u.'+key+'=?');args.append(val)
        term=request.args.get('term','').strip()
        if term: clauses.append('(u.id LIKE ? OR u.name LIKE ?)');args.extend(['%'+term+'%']*2)
        query='SELECT r.*,u.name,u.g,u.c,u.n,u.home FROM records r JOIN users u ON u.id=r.sid'
        if clauses: query+=' WHERE '+' AND '.join(clauses)
        return [dict(r) for r in conn.execute(query+' ORDER BY r.date DESC,r.id DESC',args)]
    def report_data(conn):
        rows=filtered(conn)
        if request.args.get('repeat')!='1': return rows
        try: threshold=int(request.args.get('threshold','5'))
        except ValueError: raise Problem('기준 횟수를 확인하세요.')
        if not 1<=threshold<=100000: raise Problem('기준 횟수는 1 이상으로 입력하세요.')
        grouped={}
        for r in rows:
            if r['sid'] not in grouped: grouped[r['sid']]={**r,'count':0}
            grouped[r['sid']]['count']+=1
        return sorted([r for r in grouped.values() if r['count']>=threshold],key=lambda x:(-x['count'],x['sid']))
    @app.get(PREFIX+'/api/report')
    @auth('admin')
    def report_api():
        with db() as conn: return jsonify(rows=report_data(conn))
    def xlsx(rows,name):
        wb=Workbook();ws=wb.active;ws.title='자료'
        for row in rows: ws.append(row)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type='s'
        for cell in ws[1]: cell.font=__import__('openpyxl').styles.Font(bold=True)
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(55,max(12,max(len(str(c.value or '')) for c in col)*1.7))
        ws.freeze_panes='A2';out=io.BytesIO();wb.save(out);out.seek(0)
        return send_file(out,as_attachment=True,download_name=name,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    @app.get(PREFIX+'/api/report.xlsx')
    @auth('admin')
    def export_report():
        with db() as conn: rows=report_data(conn)
        if request.args.get('repeat')=='1':
            data=[['학번','학년','반','번호','이름','담임교사','지도횟수','최근지도일','최근지도항목']]+[[r[k] for k in ('sid','g','c','n','name','home','count','date','item')] for r in rows]
        else: data=[['학번','학년','반','번호','이름','일자','점검항목','안내내용','지도교사','처리상태']]+[[r[k] for k in ('sid','g','c','n','name','date','item','note','teacher','status')] for r in rows]
        return xlsx(data,'단정체크_보고서.xlsx')
    @app.get(PREFIX+'/api/template/<kind>.xlsx')
    @auth('admin')
    def template(kind):
        if kind=='students': data=[['학년','반','번호','이름','담임교사','초기비밀번호'],[1,1,1,'김하늘','김은혜','1234']]
        elif kind=='teachers': data=[['교사아이디','이름','역할','담당학년','담당반','초기비밀번호'],['t101','김은혜','teacher',1,1,'1234']]
        else: raise Problem('양식을 확인하세요.',404)
        return xlsx(data,'단정체크_'+kind+'_양식.xlsx')
    @app.post(PREFIX+'/api/import/<kind>')
    @auth('admin')
    def import_users(kind):
        if kind not in ('students','teachers'): raise Problem('업로드 종류를 확인하세요.')
        f=request.files.get('file')
        if not f: raise Problem('Excel 파일을 선택하세요.')
        if not f.filename.lower().endswith(('.xlsx','.csv')): raise Problem('xlsx 또는 UTF-8 CSV 파일을 선택하세요. 구형 xls는 xlsx로 저장하세요.')
        try:
            if f.filename.lower().endswith('.csv'): matrix=list(csv.reader(io.StringIO(f.read().decode('utf-8-sig'))))
            else:
                raw=f.read()
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    if sum(x.file_size for x in archive.infolist())>50*1024*1024: raise Problem('압축 해제된 Excel 자료가 너무 큽니다.')
                wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
                matrix=[]
                for row in wb.active.values:
                    matrix.append(row)
                    if len(matrix)>10001: break
                wb.close()
            if not matrix or len(matrix)>10001: raise Problem('파일은 최대 10,000행까지 등록할 수 있습니다.')
        except Problem: raise
        except Exception: raise Problem('Excel 파일 형식을 확인하세요.')
        headings=[str(x or '').strip() for x in matrix[0]];errors=[];prepared=[];seen=set()
        with db() as conn:
            for number,row in enumerate(matrix[1:],2):
                if not any(x is not None and str(x).strip() for x in row): continue
                d=dict(zip(headings,[str(x) if x is not None else '' for x in row]))
                try:
                    if kind=='students':
                        udata={'role':'student','g':d.get('학년',''),'c':d.get('반',''),'n':d.get('번호',''),'name':d.get('이름',''),'home':d.get('담임교사','')}
                        uid=student_id(udata['g'],udata['c'],udata['n'])
                    else:
                        uid=d.get('교사아이디','');role={'교사':'teacher','관리자':'admin'}.get(d.get('역할'),d.get('역할') or 'teacher')
                        if role not in ('teacher','admin'): raise Problem('교사 역할을 확인하세요.')
                        udata={'id':uid,'role':role,'name':d.get('이름',''),'g':d.get('담당학년',''),'c':d.get('담당반',''),'n':'','home':''}
                    old=conn.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone()
                    if old and (old['role']=='student')!=(kind=='students'): raise Problem('다른 계정 유형과 아이디가 겹칩니다.')
                    if uid in seen: raise Problem('파일 안에 중복 아이디가 있습니다.')
                    seen.add(uid)
                    if d.get('초기비밀번호'): udata['password']=d['초기비밀번호']
                    prepared.append((make_user(udata,old),bool(old)))
                except Problem as ex: errors.append({'row':number,'error':ex.message})
            if errors: return jsonify(error='오류 행을 수정해 다시 업로드하세요. 이번 파일은 저장되지 않았습니다.',errors=errors),400
            created=updated=0
            for u,old in prepared:
                write_user(conn,u);conn.execute('DELETE FROM sessions WHERE uid=?',(u['id'],));created+=not old;updated+=old
            audit(conn,'excel_import',after={'kind':kind,'created':created,'updated':updated})
        return jsonify(message=f'신규 {created}건 / 수정 {updated}건',created=created,updated=updated)

    def snapshot(conn):
        return {'format':'danjeong-server','version':1,'created_at':stamp(),**{table:[dict(r) for r in conn.execute('SELECT * FROM '+table)] for table in ('users','items','rules','records','audit','resets','settings')}}
    def automatic_backup(conn,label):
        directory=app.config['DATA']/'backups';directory.mkdir(exist_ok=True)
        path=directory/(datetime.now(KST).strftime('%Y%m%d-%H%M%S')+'-'+label+'-'+secrets.token_hex(3)+'.json')
        path.write_text(json.dumps(snapshot(conn),ensure_ascii=False,indent=2),encoding='utf-8')
    @app.get(PREFIX+'/api/backup')
    @auth('admin')
    def backup():
        with db() as conn: data=snapshot(conn)
        return send_file(io.BytesIO(json.dumps(data,ensure_ascii=False,indent=2).encode()),as_attachment=True,download_name='단정체크_전체백업_'+today()+'.json',mimetype='application/json')
    def validate_backup(d):
        if not isinstance(d,dict) or not all(isinstance(d.get(k),list) for k in ('users','items','records')): raise Problem('올바른 단정체크 백업 파일이 아닙니다.')
        legacy=d.get('format')!='danjeong-server'
        if not legacy and d.get('version')!=1: raise Problem('지원하지 않는 백업 버전입니다.')
        result={};seen=set()
        for u in d['users']:
            uid=text(u.get('id',''),50);name=text(u.get('name',''),100);role=u.get('role')
            if not re.fullmatch(r'[A-Za-z0-9_-]{1,50}',uid) or not name or role not in ('student','teacher','admin') or uid in seen: raise Problem('백업 계정 정보가 올바르지 않습니다.')
            seen.add(uid)
            hashed=generate_password_hash(password(u.get('pw'))) if legacy else text(u.get('password_hash',''),500)
            if not hashed or (not legacy and not re.fullmatch(r'(scrypt:32768:8:1|pbkdf2:sha256:600000|pbkdf2:sha256:1000000)\$[A-Za-z0-9]{8,32}\$[a-f0-9]{64,128}',hashed)): raise Problem('백업 비밀번호 형식이 올바르지 않습니다.')
            result.setdefault('users',[]).append(dict(id=uid,name=name,role=role,password_hash=hashed,**{k:text(u.get(k,''),100) for k in ('g','c','n','home')}))
        if not any(u['role']=='admin' for u in result.get('users',[])): raise Problem('백업에는 관리자 계정이 필요합니다.')
        for table,defaults in [('items',DEFAULT_ITEMS),('rules',DEFAULT_RULES)]:
            result[table]=[]
            for i,v in enumerate(d.get(table,defaults),1):
                val=text(v if isinstance(v,str) else v.get('value',''),500)
                if not val: raise Problem('백업 항목이 비어 있습니다.')
                result[table].append({'id':i if isinstance(v,str) else int(v['id']),'value':val})
        result['records']=[]
        for r in d['records']:
            student_ids={u['id'] for u in result['users'] if u['role']=='student'}
            if r.get('sid') not in student_ids or r.get('status') not in ('pending','review','done'): raise Problem('백업 지도 기록을 확인하세요.')
            result['records'].append(dict(id=int(r['id']),sid=r['sid'],item=text(r.get('item',''),500),note=text(r.get('note',''),4000),status=r['status'],teacher_id=text(r.get('teacher_id','legacy'),50),teacher=text(r.get('teacher',''),100),date=valid_date(r.get('date')),updated_at=text(r.get('updated_at',stamp()),100),version=int(r.get('version',1))))
        result['audit']=[] if legacy else d.get('audit',[])
        result['resets']=([{'id':i,'sid':text(r.get('sid',''),50),'name':text(r.get('name',''),100),'processor':text(r.get('processor',''),100),'date':text(r.get('date',''),100)} for i,r in enumerate(d.get('reset',[]),1)] if legacy else d.get('resets',[]))
        if legacy: result['settings']=[{'key':'defaultStudentPw','value':password(d.get('defaultStudentPw','1234'))}]
        else:
            values={r['key']:r['value'] for r in d.get('settings',[])}
            result['settings']=[{'key':'defaultStudentPw','value':password(values.get('defaultStudentPw','1234'))}]
        return result
    @app.post(PREFIX+'/api/restore')
    @auth('admin')
    def restore():
        f=request.files.get('file')
        if not f: raise Problem('백업 JSON 파일을 선택하세요.')
        try: result=validate_backup(json.load(f.stream))
        except Problem: raise
        except Exception: raise Problem('백업 파일 구조를 확인하세요.')
        tables={'users':('id','name','role','password_hash','g','c','n','home'),'items':('id','value'),'rules':('id','value'),'records':('id','sid','item','note','status','teacher_id','teacher','date','updated_at','version'),'audit':('id','record_id','event','actor','at','before_json','after_json'),'resets':('id','sid','name','processor','date'),'settings':('key','value')}
        try:
            with db() as conn:
                automatic_backup(conn,'before-restore')
                for table in ('sessions','records','users','items','rules','audit','resets','settings','login_attempts'): conn.execute('DELETE FROM '+table)
                for table,keys in tables.items():
                    for row in result[table]: conn.execute('INSERT INTO '+table+'('+','.join(keys)+') VALUES('+','.join(['?']*len(keys))+')',[row[k] for k in keys])
                audit(conn,'backup_restore')
        except (KeyError,ValueError,TypeError,sqlite3.IntegrityError): raise Problem('백업 데이터가 올바르지 않아 복원하지 않았습니다.')
        return jsonify(message='복원 완료. 백업에 있는 계정으로 다시 로그인하세요.')
    @app.post(PREFIX+'/api/year-reset')
    @auth('admin')
    def year_reset():
        d=body()
        with db() as conn:
            row=conn.execute('SELECT * FROM users WHERE id=?',(g.user['id'],)).fetchone()
            if d.get('confirm')!='학년도 초기화' or not isinstance(d.get('password'),str) or not check_password_hash(row['password_hash'],d['password']): raise Problem('확인 문구와 관리자 비밀번호를 확인하세요.',403)
            automatic_backup(conn,'before-year-reset')
            for table in ('records','audit','resets','items','rules'): conn.execute('DELETE FROM '+table)
            conn.execute('DELETE FROM users WHERE id<>?',(g.user['id'],))
            conn.executemany('INSERT INTO items(value) VALUES(?)',[(x,) for x in DEFAULT_ITEMS]);conn.executemany('INSERT INTO rules(value) VALUES(?)',[(x,) for x in DEFAULT_RULES]);audit(conn,'year_reset')
        return jsonify(message='학년도 초기화 완료. 초기화 전 자동 백업을 저장했습니다.')
    return app


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--port',type=int,default=8787);parser.add_argument('--data-dir');args=parser.parse_args()
    app=create_app(args.data_dir)
    logs=app.config['DATA']/'logs';logs.mkdir(exist_ok=True)
    from logging.handlers import RotatingFileHandler
    handler=RotatingFileHandler(logs/'server.log',maxBytes=2_000_000,backupCount=3,encoding='utf-8');handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'));app.logger.addHandler(handler)
    setup=app.config['DATA']/'setup-code.txt'
    if setup.exists(): print('Initial setup code file:',setup,flush=True)
    print(f'DanjeongCheck running on http://127.0.0.1:{args.port}/danjeong/',flush=True)
    from waitress import serve
    serve(app,host='127.0.0.1',port=args.port,threads=8,max_request_body_size=12*1024*1024)

if __name__=='__main__': main()
